from datetime import date, datetime, timedelta
from decimal import Decimal
import io
import openpyxl

from django.test import TestCase
from django.utils import timezone
from rest_framework.test import APIRequestFactory, force_authenticate

from apps.products.models import Brand, Category, Product, ProductBatch, ProductUnitMeasurement
from apps.reports.services.report_builder import ReportBuilderService
from apps.reports.views.report_builder_view import (
    ReportBuilderExportAPIView,
    ReportBuilderGenerateAPIView,
)
from apps.sales.models import BankCard, Payment, Sale, SaleItem, SaleReturn, SaleReturnItem
from apps.store.models import Store, StoreUser
from apps.users.models.customers import Customer
from apps.users.models.role import Role
from apps.users.models.user import User


class SalesViewsReportTests(TestCase):
    """
    Comprehensive test suite for Sales Report:
    - Cheklar View (1 Sale = 1 row)
    - Mahsulotlar View (1 SaleItem = 1 row)
    - Dual Sheet Excel Export (Cheklar & Mahsulotlar)
    - Filtering, Server-side pagination, Sorting, RBAC & Store Isolation
    - Financial & return correctness, discount distribution, no double counting
    """

    @classmethod
    def setUpTestData(cls):
        # Stores
        cls.store1 = Store.objects.create(name="Store Alpha", address="Alpha Address", phone_number="+998901111111")
        cls.store2 = Store.objects.create(name="Store Beta", address="Beta Address", phone_number="+998902222222")

        # Roles
        cls.role_viewer = Role.objects.create(name="Sales Viewer", permissions=["reports.view", "reports.sales.view"])
        cls.role_exporter = Role.objects.create(name="Sales Exporter", permissions=["reports.view", "reports.sales.view", "reports.sales.export"])
        cls.role_no_view = Role.objects.create(name="No View Role", permissions=["reports.view"])

        # Users
        cls.admin = User.objects.create(phone_number="+998900000001", full_name="Super Admin", is_superuser=True, is_staff=True)
        cls.seller1 = User.objects.create(phone_number="+998900000002", full_name="Seller Sam", is_superuser=False)
        cls.seller2 = User.objects.create(phone_number="+998900000003", full_name="Seller Sally", is_superuser=False)
        cls.store1_mgr = User.objects.create(phone_number="+998900000004", full_name="Store 1 Mgr", role=cls.role_viewer, is_superuser=False)
        cls.store2_mgr = User.objects.create(phone_number="+998900000005", full_name="Store 2 Mgr", role=cls.role_viewer, is_superuser=False)
        cls.export_user = User.objects.create(phone_number="+998900000006", full_name="Export User", role=cls.role_exporter, is_superuser=False)
        cls.no_perm_user = User.objects.create(phone_number="+998900000007", full_name="No Perm User", role=cls.role_no_view, is_superuser=False)

        StoreUser.objects.create(user=cls.store1_mgr, store=cls.store1, is_active=True)
        StoreUser.objects.create(user=cls.store2_mgr, store=cls.store2, is_active=True)
        StoreUser.objects.create(user=cls.export_user, store=cls.store1, is_active=True)
        StoreUser.objects.create(user=cls.no_perm_user, store=cls.store1, is_active=True)
        StoreUser.objects.create(user=cls.seller1, store=cls.store1, is_active=True)
        StoreUser.objects.create(user=cls.seller2, store=cls.store2, is_active=True)
        StoreUser.objects.create(user=cls.admin, store=cls.store1, is_active=True)
        StoreUser.objects.create(user=cls.admin, store=cls.store2, is_active=True)

        # Categories, Brands & Units
        cls.cat_brakes = Category.objects.create(name="Tormoz tizimi")
        cls.cat_filters = Category.objects.create(name="Filtrlar")
        cls.brand_bosch = Brand.objects.create(name="Bosch")
        cls.brand_mann = Brand.objects.create(name="Mann")
        cls.unit_dona = ProductUnitMeasurement.objects.create(measurement="Dona")
        cls.unit_juft = ProductUnitMeasurement.objects.create(measurement="Juft")

        # Products
        cls.prod_pads = Product.objects.create(
            name="Kolodka Bosch",
            barcode="4781001000018",
            sku="BP-001",
            category=cls.cat_brakes,
            brand=cls.brand_bosch,
            unit_measurement=cls.unit_juft,
            status=Product.ProductStatus.ACTIVE,
        )
        cls.prod_oil_filter = Product.objects.create(
            name="Moy filtri Mann",
            barcode="4781001000025",
            sku="OF-002",
            category=cls.cat_filters,
            brand=cls.brand_mann,
            unit_measurement=cls.unit_dona,
            status=Product.ProductStatus.ACTIVE,
        )
        cls.prod_air_filter = Product.objects.create(
            name="Havo filtri Mann",
            barcode="4781001000032",
            sku="AF-003",
            category=cls.cat_filters,
            brand=cls.brand_mann,
            unit_measurement=cls.unit_dona,
            status=Product.ProductStatus.ACTIVE,
        )

        cls.customer1 = Customer.objects.create(full_name="Akmal Saidov", phone_number="+998901234567")
        cls.customer2 = Customer.objects.create(full_name="Botir Zokirov", phone_number="+998909876543")

        cls.factory = APIRequestFactory()

    def _create_sale(
        self,
        store=None,
        seller=None,
        customer=None,
        items=None,
        discount_amount=Decimal("0.00"),
        status=Sale.Status.PAID,
        payment_type=Sale.PaymentType.CASH,
        created_at=None,
    ):
        store = store or self.store1
        seller = seller or self.seller1
        items_data = items or []

        gross_total = sum(i["quantity"] * i["unit_price"] for i in items_data)
        net_total = max(Decimal("0.00"), gross_total - discount_amount)
        paid_amount = net_total if status == Sale.Status.PAID else Decimal("0.00")

        sale = Sale.objects.create(
            store=store,
            seller=seller,
            customer=customer,
            total_amount=net_total,
            paid_amount=paid_amount,
            status=status,
            payment_type=payment_type,
            discount_amount=discount_amount,
            discount_type=Sale.DiscountType.FIXED if discount_amount > 0 else None,
            discount_value=discount_amount,
        )
        if created_at:
            Sale.objects.filter(id=sale.id).update(created_at=created_at)
            sale.refresh_from_db()

        created_items = []
        for i in items_data:
            si = SaleItem.objects.create(
                sale=sale,
                product=i["product"],
                quantity=i["quantity"],
                purchase_price=i.get("purchase_price", Decimal("50.00")),
                unit_price=i["unit_price"],
                total_price=i["quantity"] * i["unit_price"],
                returned_quantity=i.get("returned_quantity", Decimal("0.00")),
            )
            created_items.append(si)

        return sale, created_items

    # ─────────────────────────────────────────────────────────────
    # 1. CHEKLAR VIEW TESTS
    # ─────────────────────────────────────────────────────────────
    def test_01_cheklar_single_sale_produces_single_row(self):
        sale, items = self._create_sale(
            items=[
                {"product": self.prod_pads, "quantity": Decimal("2"), "unit_price": Decimal("100.00")},
                {"product": self.prod_oil_filter, "quantity": Decimal("3"), "unit_price": Decimal("50.00")},
            ]
        )
        res = ReportBuilderService.generate({"report_type": "sales", "view": "receipts"}, self.admin)
        self.assertEqual(res["total"], 1)
        self.assertEqual(len(res["rows"]), 1)
        row = res["rows"][0]
        self.assertEqual(row["id"], sale.id)
        self.assertEqual(row["total"], "350.00")
        self.assertEqual(row["store"], self.store1.name)

    def test_02_cheklar_multiple_sales_produce_multiple_rows(self):
        self._create_sale(items=[{"product": self.prod_pads, "quantity": Decimal("1"), "unit_price": Decimal("100.00")}])
        self._create_sale(items=[{"product": self.prod_oil_filter, "quantity": Decimal("2"), "unit_price": Decimal("50.00")}])
        res = ReportBuilderService.generate({"report_type": "sales", "view": "receipts"}, self.admin)
        self.assertEqual(res["total"], 2)
        self.assertEqual(len(res["rows"]), 2)

    def test_03_cheklar_existing_filters_work(self):
        self._create_sale(
            payment_type=Sale.PaymentType.CASH,
            status=Sale.Status.PAID,
            items=[{"product": self.prod_pads, "quantity": Decimal("1"), "unit_price": Decimal("100.00")}],
        )
        self._create_sale(
            payment_type=Sale.PaymentType.CARD,
            status=Sale.Status.DEBT,
            items=[{"product": self.prod_oil_filter, "quantity": Decimal("1"), "unit_price": Decimal("60.00")}],
        )

        res_cash = ReportBuilderService.generate({"report_type": "sales", "payment_type": "cash"}, self.admin)
        self.assertEqual(res_cash["total"], 1)
        self.assertEqual(res_cash["rows"][0]["payment"], "Naqd")

        res_debt = ReportBuilderService.generate({"report_type": "sales", "status": "debt"}, self.admin)
        self.assertEqual(res_debt["total"], 1)
        self.assertEqual(res_debt["rows"][0]["status"], "Qarz")

    def test_04_cheklar_rbac_permissions(self):
        self._create_sale(items=[{"product": self.prod_pads, "quantity": Decimal("1"), "unit_price": Decimal("100.00")}])
        req = self.factory.get("/api/reports/builder/?report_type=sales")
        force_authenticate(req, user=self.no_perm_user)
        view = ReportBuilderGenerateAPIView.as_view()
        response = view(req)
        self.assertEqual(response.status_code, 403)

        force_authenticate(req, user=self.store1_mgr)
        response = view(req)
        self.assertEqual(response.status_code, 200)

    def test_05_cheklar_store_isolation(self):
        self._create_sale(store=self.store1, items=[{"product": self.prod_pads, "quantity": Decimal("1"), "unit_price": Decimal("100.00")}])
        self._create_sale(store=self.store2, items=[{"product": self.prod_oil_filter, "quantity": Decimal("1"), "unit_price": Decimal("50.00")}])

        req = self.factory.get("/api/reports/builder/?report_type=sales")
        force_authenticate(req, user=self.store1_mgr)
        response = ReportBuilderGenerateAPIView.as_view()(req)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data["total"], 1)
        self.assertEqual(response.data["rows"][0]["store"], self.store1.name)

    # ─────────────────────────────────────────────────────────────
    # 2. MAHSULOTLAR VIEW TESTS
    # ─────────────────────────────────────────────────────────────
    def test_06_mahsulotlar_1_sale_with_3_items_produces_3_rows(self):
        sale, items = self._create_sale(
            customer=self.customer1,
            seller=self.seller1,
            items=[
                {"product": self.prod_pads, "quantity": Decimal("2"), "unit_price": Decimal("100.00")},
                {"product": self.prod_oil_filter, "quantity": Decimal("1"), "unit_price": Decimal("45.00")},
                {"product": self.prod_air_filter, "quantity": Decimal("3"), "unit_price": Decimal("60.00")},
            ],
        )
        res = ReportBuilderService.generate({"report_type": "sales", "view": "items"}, self.admin)
        self.assertEqual(res["total"], 3)
        self.assertEqual(len(res["rows"]), 3)

    def test_07_mahsulotlar_different_products_remain_separate(self):
        self._create_sale(
            items=[
                {"product": self.prod_pads, "quantity": Decimal("2"), "unit_price": Decimal("100.00")},
                {"product": self.prod_oil_filter, "quantity": Decimal("5"), "unit_price": Decimal("40.00")},
            ]
        )
        res = ReportBuilderService.generate({"report_type": "sales", "view": "items"}, self.admin)
        products = {r["product"] for r in res["rows"]}
        self.assertEqual(products, {"Kolodka Bosch", "Moy filtri Mann"})

    def test_08_to_15_mahsulotlar_row_fields_accurate(self):
        t_now = timezone.now()
        sale, items = self._create_sale(
            customer=self.customer1,
            seller=self.seller1,
            created_at=t_now,
            items=[
                {"product": self.prod_pads, "quantity": Decimal("2.50"), "unit_price": Decimal("120.00")},
            ],
        )
        res = ReportBuilderService.generate({"report_type": "sales", "view": "items"}, self.admin)
        row = res["rows"][0]

        # 08: quantity
        self.assertEqual(row["quantity"], "2.50")
        self.assertEqual(row["net_quantity"], "2.50")
        # 09: product name
        self.assertEqual(row["product"], "Kolodka Bosch")
        # 10: SKU
        self.assertEqual(row["sku"], "BP-001")
        # 11: barcode
        self.assertEqual(row["barcode"], "4781001000018")
        # 12: date
        expected_date = timezone.localtime(t_now).strftime("%d.%m.%Y %H:%M")
        self.assertEqual(row["date"], expected_date)
        # 13: check number
        self.assertEqual(row["sale_id"], sale.id)
        # 14: seller
        self.assertEqual(row["seller"], "Seller Sam")
        # 15: customer
        self.assertEqual(row["customer"], "Akmal Saidov")
        # additional metadata
        self.assertEqual(row["category"], "Tormoz tizimi")
        self.assertEqual(row["brand"], "Bosch")
        self.assertEqual(row["unit"], "Juft")

    def test_16_mahsulotlar_filters_work(self):
        s1, _ = self._create_sale(
            seller=self.seller1,
            customer=self.customer1,
            items=[{"product": self.prod_pads, "quantity": Decimal("1"), "unit_price": Decimal("100.00")}],
        )
        s2, _ = self._create_sale(
            seller=self.seller2,
            customer=self.customer2,
            items=[{"product": self.prod_oil_filter, "quantity": Decimal("2"), "unit_price": Decimal("50.00")}],
        )

        res_seller = ReportBuilderService.generate({"report_type": "sales", "view": "items", "seller_id": str(self.seller1.id)}, self.admin)
        self.assertEqual(res_seller["total"], 1)
        self.assertEqual(res_seller["rows"][0]["product"], "Kolodka Bosch")

        res_search_prod = ReportBuilderService.generate({"report_type": "sales", "view": "items", "search": "OF-002"}, self.admin)
        self.assertEqual(res_search_prod["total"], 1)
        self.assertEqual(res_search_prod["rows"][0]["sku"], "OF-002")

    def test_17_mahsulotlar_server_side_sorting(self):
        self._create_sale(items=[
            {"product": self.prod_pads, "quantity": Decimal("10"), "unit_price": Decimal("10.00")},    # total 100
            {"product": self.prod_oil_filter, "quantity": Decimal("1"), "unit_price": Decimal("500.00")}, # total 500
        ])
        res_sort_total_desc = ReportBuilderService.generate(
            {"report_type": "sales", "view": "items", "sort_by": "total", "order": "desc"}, self.admin
        )
        self.assertEqual(res_sort_total_desc["rows"][0]["product"], "Moy filtri Mann")
        self.assertEqual(res_sort_total_desc["rows"][1]["product"], "Kolodka Bosch")

        res_sort_qty_desc = ReportBuilderService.generate(
            {"report_type": "sales", "view": "items", "sort_by": "quantity", "order": "desc"}, self.admin
        )
        self.assertEqual(res_sort_qty_desc["rows"][0]["product"], "Kolodka Bosch")

    def test_18_mahsulotlar_server_side_pagination(self):
        for i in range(15):
            self._create_sale(items=[{"product": self.prod_pads, "quantity": Decimal("1"), "unit_price": Decimal("10.00")}])

        res_page1 = ReportBuilderService.generate({"report_type": "sales", "view": "items", "page": "1", "limit": "5"}, self.admin)
        self.assertEqual(res_page1["total"], 15)
        self.assertEqual(len(res_page1["rows"]), 5)
        self.assertEqual(res_page1["page"], 1)
        self.assertEqual(res_page1["limit"], 5)

        res_page2 = ReportBuilderService.generate({"report_type": "sales", "view": "items", "page": "2", "limit": "5"}, self.admin)
        self.assertEqual(len(res_page2["rows"]), 5)
        self.assertEqual(res_page2["page"], 2)

    # ─────────────────────────────────────────────────────────────
    # 3. EXCEL EXPORT (2 SHEETS) TESTS
    # ─────────────────────────────────────────────────────────────
    def test_19_to_25_excel_dual_sheets_verification(self):
        self._create_sale(
            customer=self.customer1,
            seller=self.seller1,
            items=[
                {"product": self.prod_pads, "quantity": Decimal("2"), "unit_price": Decimal("100.00")},
                {"product": self.prod_oil_filter, "quantity": Decimal("3"), "unit_price": Decimal("50.00")},
            ],
        )

        req = self.factory.get("/api/reports/builder/export/?report_type=sales&export_type=excel")
        force_authenticate(req, user=self.export_user)
        response = ReportBuilderExportAPIView.as_view()(req)

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        self.assertIn("Sales_Report_", response["Content-Disposition"])

        wb = openpyxl.load_workbook(io.BytesIO(response.content))
        # 19: exactly 2 sheets
        sheet_names = wb.sheetnames
        self.assertEqual(len(sheet_names), 2)
        self.assertEqual(sheet_names, ["Cheklar", "Mahsulotlar"])

        ws_cheklar = wb["Cheklar"]
        ws_items = wb["Mahsulotlar"]

        # 20: Cheklar has 1 data row (header row 1 + data row 2 = 2 rows)
        self.assertEqual(ws_cheklar.max_row, 2)
        # 21: Mahsulotlar has 2 data rows (header row 1 + 2 items = 3 rows)
        self.assertEqual(ws_items.max_row, 3)

        # 23: Excel headers correct
        cheklar_headers = [cell.value for cell in ws_cheklar[1]]
        self.assertEqual(cheklar_headers[:5], ["Chek №", "Sana", "Do'kon", "Mijoz", "Sotuvchi"])
        items_headers = [cell.value for cell in ws_items[1]]
        self.assertEqual(items_headers[:5], ["Check №", "Sana", "Do'kon", "Tovar", "SKU"])

        # 25: Valid Excel table exists on each sheet
        self.assertIn("CheklarTable", ws_cheklar.tables)
        self.assertIn("MahsulotlarTable", ws_items.tables)

    # ─────────────────────────────────────────────────────────────
    # 4. BUSINESS LOGIC & FINANCIAL CONSISTENCY
    # ─────────────────────────────────────────────────────────────
    def test_26_proportional_discount_distribution(self):
        # Item 1: 2 * 100 = 200
        # Item 2: 2 * 50 = 100
        # Subtotal = 300, Discount = 30 (10%)
        # Total = 270
        sale, items = self._create_sale(
            discount_amount=Decimal("30.00"),
            items=[
                {"product": self.prod_pads, "quantity": Decimal("2"), "unit_price": Decimal("100.00")},
                {"product": self.prod_oil_filter, "quantity": Decimal("2"), "unit_price": Decimal("50.00")},
            ],
        )
        res = ReportBuilderService.generate({"report_type": "sales", "view": "items"}, self.admin)
        rows_by_prod = {r["product"]: r for r in res["rows"]}

        row1 = rows_by_prod["Kolodka Bosch"]
        self.assertEqual(row1["discount"], "20.00")
        self.assertEqual(row1["total"], "180.00")

        row2 = rows_by_prod["Moy filtri Mann"]
        self.assertEqual(row2["discount"], "10.00")
        self.assertEqual(row2["total"], "90.00")

        # Sum of item net totals equals sale net total
        sum_item_totals = Decimal(row1["total"]) + Decimal(row2["total"])
        self.assertEqual(sum_item_totals, Decimal("270.00"))

    def test_27_to_29_partial_and_full_returns(self):
        # Sale: 4 units @ 100 = 400. 1 unit returned.
        sale, items = self._create_sale(
            items=[{"product": self.prod_pads, "quantity": Decimal("4"), "unit_price": Decimal("100.00"), "returned_quantity": Decimal("1")}]
        )
        res = ReportBuilderService.generate({"report_type": "sales", "view": "items"}, self.admin)
        row = res["rows"][0]
        self.assertEqual(row["quantity"], "4")
        self.assertEqual(row["returned_quantity"], "1")
        self.assertEqual(row["net_quantity"], "3")
        self.assertEqual(row["total"], "300.00")

        # Full return sale
        sale_ret, items_ret = self._create_sale(
            status=Sale.Status.RETURNED,
            items=[{"product": self.prod_oil_filter, "quantity": Decimal("2"), "unit_price": Decimal("50.00"), "returned_quantity": Decimal("2")}],
        )
        res_ret = ReportBuilderService.generate({"report_type": "sales", "view": "items", "search": str(sale_ret.id)}, self.admin)
        row_ret = res_ret["rows"][0]
        self.assertEqual(row_ret["net_quantity"], "0")
        self.assertEqual(row_ret["total"], "0.00")

    def test_30_no_double_counting_across_views(self):
        sale, items = self._create_sale(
            discount_amount=Decimal("50.00"),
            items=[
                {"product": self.prod_pads, "quantity": Decimal("3"), "unit_price": Decimal("100.00")},       # gross 300
                {"product": self.prod_oil_filter, "quantity": Decimal("2"), "unit_price": Decimal("100.00")},  # gross 200
            ],
        )
        # Total before discount = 500, discount = 50 -> net_total = 450
        res_cheklar = ReportBuilderService.generate({"report_type": "sales", "view": "receipts"}, self.admin)
        cheklar_total = Decimal(res_cheklar["rows"][0]["total"])
        self.assertEqual(cheklar_total, Decimal("450.00"))

        res_items = ReportBuilderService.generate({"report_type": "sales", "view": "items"}, self.admin)
        items_total = sum(Decimal(r["total"]) for r in res_items["rows"])
        self.assertEqual(items_total, cheklar_total)

    # ─────────────────────────────────────────────────────────────
    # 5. PERFORMANCE & ZERO N+1 VERIFICATION
    # ─────────────────────────────────────────────────────────────
    def test_31_and_32_query_count_and_no_n_plus_one(self):
        for i in range(10):
            self._create_sale(
                customer=self.customer1,
                seller=self.seller1,
                items=[
                    {"product": self.prod_pads, "quantity": Decimal("1"), "unit_price": Decimal("100.00")},
                    {"product": self.prod_oil_filter, "quantity": Decimal("2"), "unit_price": Decimal("50.00")},
                    {"product": self.prod_air_filter, "quantity": Decimal("3"), "unit_price": Decimal("30.00")},
                ],
            )

        # In items view, generating page of 10 items should execute constant O(1) queries (count + aggregate + items slice)
        with self.assertNumQueries(3):
            res = ReportBuilderService.generate({"report_type": "sales", "view": "items", "limit": "10"}, self.admin)
            self.assertEqual(len(res["rows"]), 10)

    def test_33_csv_export_clean_format(self):
        self._create_sale(items=[{"product": self.prod_pads, "quantity": Decimal("1"), "unit_price": Decimal("100.00")}])
        req_receipts = self.factory.get("/api/reports/builder/export/?report_type=sales&export_type=csv&view=receipts")
        force_authenticate(req_receipts, user=self.export_user)
        res_receipts = ReportBuilderExportAPIView.as_view()(req_receipts)
        self.assertEqual(res_receipts.status_code, 200)
        csv_text = res_receipts.content.decode("utf-8-sig")
        lines = [line for line in csv_text.splitlines() if line.strip()]
        self.assertEqual(len(lines), 2)  # 1 header + 1 data row (clean, no trailing summary)
        self.assertIn("Chek №", lines[0])

        req_items = self.factory.get("/api/reports/builder/export/?report_type=sales&export_type=csv&view=items")
        force_authenticate(req_items, user=self.export_user)
        res_items = ReportBuilderExportAPIView.as_view()(req_items)
        self.assertEqual(res_items.status_code, 200)
        csv_items_text = res_items.content.decode("utf-8-sig")
        items_lines = [line for line in csv_items_text.splitlines() if line.strip()]
        self.assertEqual(len(items_lines), 2)  # 1 header + 1 data row
        self.assertIn("Tovar", items_lines[0])
        self.assertIn("Kolodka Bosch", items_lines[1])

    def test_34_excel_export_store_isolation(self):
        self._create_sale(store=self.store1, items=[{"product": self.prod_pads, "quantity": Decimal("1"), "unit_price": Decimal("100.00")}])
        self._create_sale(store=self.store2, items=[{"product": self.prod_oil_filter, "quantity": Decimal("2"), "unit_price": Decimal("50.00")}])

        req = self.factory.get("/api/reports/builder/export/?report_type=sales&export_type=excel")
        force_authenticate(req, user=self.export_user)  # Assigned to store1 with export role
        res = ReportBuilderExportAPIView.as_view()(req)
        self.assertEqual(res.status_code, 200)

        wb = openpyxl.load_workbook(io.BytesIO(res.content))
        ws_cheklar = wb["Cheklar"]
        ws_items = wb["Mahsulotlar"]

        # Only Store 1 data should be exported
        self.assertEqual(ws_cheklar.max_row, 2)  # 1 header + 1 sale
        self.assertEqual(ws_items.max_row, 2)    # 1 header + 1 item
        self.assertEqual(ws_cheklar.cell(row=2, column=3).value, self.store1.name)
        self.assertEqual(ws_items.cell(row=2, column=3).value, self.store1.name)

    def test_35_export_rbac_forbidden_without_permission(self):
        req = self.factory.get("/api/reports/builder/export/?report_type=sales&export_type=excel")
        force_authenticate(req, user=self.no_perm_user)
        res = ReportBuilderExportAPIView.as_view()(req)
        self.assertEqual(res.status_code, 403)

    def test_36_sorting_by_date_and_product(self):
        t1 = timezone.now() - timedelta(days=2)
        t2 = timezone.now() - timedelta(days=1)
        self._create_sale(created_at=t1, items=[{"product": self.prod_pads, "quantity": Decimal("1"), "unit_price": Decimal("100.00")}])
        self._create_sale(created_at=t2, items=[{"product": self.prod_oil_filter, "quantity": Decimal("1"), "unit_price": Decimal("50.00")}])

        # Sort by date asc
        res_date_asc = ReportBuilderService.generate(
            {"report_type": "sales", "view": "items", "sort_by": "date", "order": "asc"}, self.admin
        )
        self.assertEqual(res_date_asc["rows"][0]["product"], "Kolodka Bosch")
        self.assertEqual(res_date_asc["rows"][1]["product"], "Moy filtri Mann")

        # Sort by product name asc
        res_prod_asc = ReportBuilderService.generate(
            {"report_type": "sales", "view": "items", "sort_by": "product", "order": "asc"}, self.admin
        )
        self.assertEqual(res_prod_asc["rows"][0]["product"], "Kolodka Bosch")
        self.assertEqual(res_prod_asc["rows"][1]["product"], "Moy filtri Mann")

